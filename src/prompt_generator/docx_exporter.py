"""Export verified prompt results to structured DOCX files."""

from pathlib import Path
from typing import List

from prompt_generator.models import Scene


def export_prompts_to_docx(scenes: List[Scene], output_path: Path) -> Path:
    """Export sorted scenes to a formatted DOCX document.

    Formatting strictly adheres to:
    Scene <N>: <prompt>
    """
    try:
        import docx
    except ImportError:
        raise ImportError("python-docx is required to export DOCX files.")

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    doc = docx.Document()

    # Document Title
    title = doc.add_heading("Generated Visual Prompts", level=0)
    title.paragraph_format.space_after = docx.shared.Pt(12)

    import re

    for scene in scenes:
        p = doc.add_paragraph()
        
        # Label
        run_label = p.add_run(f"Scene {scene.scene_number}: ")
        run_label.bold = True
        run_label.font.no_proof = True
        
        # Clean redundant leading 'Scene N:' if already provided by model
        prompt_text = scene.generated_prompt.strip()
        cleaned_prompt = re.sub(rf"^(?:Scene|SCENE|scene)\s+{scene.scene_number}[:.]?\s*", "", prompt_text)
        
        # Prompt Body
        run_body = p.add_run(cleaned_prompt)
        run_body.font.no_proof = True
        
        p.paragraph_format.space_after = docx.shared.Pt(8)

    doc.save(out)
    return out
